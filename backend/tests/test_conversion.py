"""抽取模組 —— 正規格式轉換測試(每種格式 / 結構 / 邊界都實測)。

格式轉換是整條流程最容易掉東西的環節,本檔分兩層:
  A. 合成夾具(self-contained):程式即時造出 DOCX/PDF/HTML/RTF,精準斷言
     「文字 / 表格 / 圖片 的順序與位置」「超連結 / 標題 / 清單 結構」「去重」「邊界錯誤」。
     不依賴外部素材 → CI 一定能跑,且能斷言肉眼難查的「圖片排列位置」。
  B. 真實素材(input-example):有檔才跑(缺檔 skip),驗真實世界稿件。

對照舊版能力(ConvertAPI→mammoth 的 inline 保位)見 docs/EXTRACTION-MODULE.md。
"""

from __future__ import annotations

import io
import os

import pytest

from app.services.extract import extract_document
from app.services.markdown import md_to_html

EX = os.path.join(os.path.dirname(__file__), "..", "..", "input-example")


# ───────────────────────── 工具:造素材 ─────────────────────────

def _png(color: str = "red", size: int = 12) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (size, size), color).save(buf, "PNG")
    return buf.getvalue()


def _png_path(tmp_path, color: str = "red") -> str:
    p = tmp_path / f"img_{color}.png"
    p.write_bytes(_png(color))
    return str(p)


def _make_docx(path: str) -> None:
    """造一份「全功能」DOCX:標題 / 段落 / 內嵌圖(夾在兩段之間)/ 超連結 / 清單 / 表格。"""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    def add_hyperlink(paragraph, url: str, text: str) -> None:
        rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        hl.append(run)
        paragraph._p.append(hl)

    d = docx.Document()
    d.add_heading("頭條標題", level=1)
    d.add_paragraph("ALPHA_前段文字")
    pic_para = d.add_paragraph()
    pic_para.add_run().add_picture(io.BytesIO(_png("red")))
    d.add_paragraph("BRAVO_後段文字")
    link_para = d.add_paragraph()
    link_para.add_run("前綴 ")
    add_hyperlink(link_para, "https://blocktempo.ai/", "動區連結")
    d.add_paragraph("項目一", style="List Bullet")
    d.add_paragraph("項目二", style="List Bullet")
    tbl = d.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "表頭A"
    tbl.cell(0, 1).text = "表頭B"
    tbl.cell(1, 0).text = "值1"
    tbl.cell(1, 1).text = "值2"
    d.save(path)


def _make_pdf_positioned(path: str) -> None:
    """造 PDF:上方文字 → 中間圖片 → 下方文字,用座標斷言閱讀順序。"""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "ALPHA_TOP_TEXT")
    page.insert_image(fitz.Rect(72, 200, 172, 300), stream=_png("blue"))
    page.insert_text((72, 420), "BRAVO_BOTTOM_TEXT")
    doc.save(path)
    doc.close()


def _make_pdf_dup_image(path: str) -> None:
    """同一張圖放在兩頁 → 測 xref 全域去重。"""
    import fitz

    png = _png("green")
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_image(fitz.Rect(72, 100, 172, 200), stream=png)
        # ASCII(fitz 內建字型無法描繪 CJK;需足量文字避免被當成掃描頁)
        page.insert_text((72, 300), f"Page {i} body text content here for the dedup fixture.")
    doc.save(path)
    doc.close()


# ───────────────── A. DOCX 合成:順序 / 結構 / 圖片位置 ─────────────────

def test_docx_image_position_between_paragraphs(tmp_path):
    """關鍵:圖片佔位符必須落在前段與後段「之間」,不是被丟到最後。"""
    p = str(tmp_path / "full.docx")
    _make_docx(p)
    text, images = extract_document(p)
    assert len(images) == 1
    assert "{{IMG0}}" in text
    assert text.index("ALPHA_前段文字") < text.index("{{IMG0}}") < text.index("BRAVO_後段文字")


def test_docx_heading_and_list_structure(tmp_path):
    p = str(tmp_path / "full.docx")
    _make_docx(p)
    text, _ = extract_document(p)
    assert "# 頭條標題" in text                      # 標題 → markdown #
    assert "- 項目一" in text and "- 項目二" in text   # 清單 → markdown -


def test_docx_hyperlink_preserved(tmp_path):
    """超連結要轉成 [text](url) —— python-docx 預設 .text 會把 hyperlink 文字丟掉。"""
    p = str(tmp_path / "full.docx")
    _make_docx(p)
    text, _ = extract_document(p)
    assert "[動區連結](https://blocktempo.ai/)" in text


def test_docx_table_to_markdown(tmp_path):
    p = str(tmp_path / "full.docx")
    _make_docx(p)
    text, _ = extract_document(p)
    assert "| 表頭A | 表頭B |" in text
    assert "| 值1 | 值2 |" in text
    assert "| --- | --- |" in text


# ───────────────── A. PDF 合成:閱讀順序 / 去重 ─────────────────

def test_pdf_image_position_reading_order(tmp_path):
    """關鍵:PDF 圖片要依座標插在上/下文字之間(舊 bug 是全丟到頁尾)。"""
    p = str(tmp_path / "pos.pdf")
    _make_pdf_positioned(p)
    text, images = extract_document(p)
    assert len(images) == 1
    assert "{{IMG0}}" in text
    assert text.index("ALPHA_TOP_TEXT") < text.index("{{IMG0}}") < text.index("BRAVO_BOTTOM_TEXT")


def test_pdf_duplicate_image_deduped(tmp_path):
    """同圖跨兩頁 → 去重後只剩 1 張。"""
    p = str(tmp_path / "dup.pdf")
    _make_pdf_dup_image(p)
    _, images = extract_document(p)
    assert len(images) == 1


def test_pdf_borderless_table_via_pdfplumber(tmp_path):
    """無框線表格(純文字排版,無線條)→ PyMuPDF find_tables 抓不到 → pdfplumber 補。"""
    import fitz

    # 用 ASCII(fitz 內建字型無法描繪 CJK 字形,與本測試的 fallback 邏輯無關)
    doc = fitz.open()
    page = doc.new_page()
    cols = [72, 220, 360]
    grid = [["Item", "Qty", "City"], ["Apple", "30", "Taipei"], ["Banana", "25", "Kaohsiung"]]
    for ri, row in enumerate(grid):
        for ci, cell in enumerate(row):
            page.insert_text((cols[ci], 100 + ri * 30), cell)
    p = str(tmp_path / "borderless.pdf")
    doc.save(p)
    doc.close()
    # fitz 單獨應抓不到(無線條);整體 extract 應靠 pdfplumber fallback 補出 markdown 表
    assert len(fitz.open(p)[0].find_tables().tables) == 0
    text, _ = extract_document(p)
    assert "| --- |" in text and "Qty" in text and "Kaohsiung" in text


def test_pdf_scanned_ocr(tmp_path):
    """掃描/圖片型 PDF(無可抽文字)→ RapidOCR(PP-OCRv5)補文字(需 ocr extra)。"""
    pytest.importorskip("rapidocr")
    import io

    from PIL import Image, ImageDraw

    import fitz

    img = Image.new("RGB", (820, 200), "white")
    ImageDraw.Draw(img).text((30, 80), "SCANNED OCR TEST CONTENT LINE 2026", fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = fitz.open()
    page = doc.new_page(width=820, height=200)
    page.insert_image(fitz.Rect(0, 0, 820, 200), stream=buf.getvalue())
    p = str(tmp_path / "scan.pdf")
    doc.save(p)
    doc.close()
    text, _ = extract_document(p)
    assert "SCANNEDOCRTESTCONTENTLINE2026" in text.replace(" ", "")  # OCR 可能不還原字間空白


def test_ocr_html_table_helper():
    """RapidTable 輸出的 HTML 表 → markdown(解析器 + 結構閘,確定性、不依賴模型)。"""
    from app.services.extract import _html_table_rows, _rows_to_md

    html = ("<html><body><table>"
            "<tr><td>Item</td><td>Qty</td></tr>"
            "<tr><td>Apple</td><td>30</td></tr>"
            "<tr><td></td><td></td></tr></table></body></html>")
    rows = _html_table_rows(html)
    assert ["Item", "Qty"] in rows
    md = _rows_to_md(rows)
    assert "| Item | Qty |" in md and "| Apple | 30 |" in md


def test_scanned_table_to_markdown(tmp_path):
    """端到端:掃描表格影像 PDF → OCR(PP-OCRv5)+ RapidTable → markdown 表格。"""
    pytest.importorskip("rapidocr")
    pytest.importorskip("rapid_table")
    import io

    from PIL import Image, ImageDraw

    import fitz

    img = Image.new("RGB", (400, 160), "white")
    d = ImageDraw.Draw(img)
    for x in (10, 140, 270, 390):
        d.line([(x, 10), (x, 150)], fill="black")
    for y in (10, 50, 90, 130, 150):
        d.line([(10, y), (390, y)], fill="black")
    cells = [["Item", "Qty", "City"], ["Apple", "30", "Taipei"], ["Banana", "25", "Osaka"]]
    for r, row in enumerate(cells):
        for c, t in enumerate(row):
            d.text((20 + c * 130, 20 + r * 40), t, fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = fitz.open()
    page = doc.new_page(width=400, height=160)
    page.insert_image(fitz.Rect(0, 0, 400, 160), stream=buf.getvalue())
    p = str(tmp_path / "scan_table.pdf")
    doc.save(p)
    doc.close()
    text, _ = extract_document(p)
    # 結構化為 markdown 表(小圖 OCR 可能個別字噪訊,驗結構 + 數個可靠儲存格)
    assert "| --- |" in text
    assert "Apple" in text and "Banana" in text and "Taipei" in text


def test_pdf_scanned_without_ocr_raises(tmp_path, monkeypatch):
    """掃描 PDF 但 OCR 不可用/失敗 → 明確報錯,不讓空白內容靜默跑下游 AI 流程。"""
    import io

    from PIL import Image, ImageDraw

    import fitz

    from app.services import extract as ex

    monkeypatch.setattr(ex, "_ocr_page", lambda page: "")     # OCR 救不回
    monkeypatch.setattr(ex, "_ocr_engine_get", lambda: None)  # 模擬未安裝

    img = Image.new("RGB", (640, 200), "white")
    ImageDraw.Draw(img).text((30, 80), "UNREADABLE SCAN", fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = fitz.open()
    page = doc.new_page(width=640, height=200)
    page.insert_image(fitz.Rect(0, 0, 640, 200), stream=buf.getvalue())
    p = str(tmp_path / "scan_noocr.pdf")
    doc.save(p)
    doc.close()
    with pytest.raises(ValueError, match="掃描"):
        extract_document(p)


def test_pdf_multipage_scan_with_thin_text_layer_raises(tmp_path, monkeypatch):
    """多頁掃描檔常帶薄文字層(頁碼/頁尾)讓全文勉強過字數門檻 → 仍須被佔比守門擋下。"""
    import io

    from PIL import Image, ImageDraw

    import fitz

    from app.services import extract as ex

    monkeypatch.setattr(ex, "_ocr_page", lambda page: "")
    monkeypatch.setattr(ex, "_ocr_engine_get", lambda: None)

    scan = Image.new("RGB", (560, 720), "white")
    ImageDraw.Draw(scan).text((40, 300), "scanned body (unreadable)", fill="black")
    buf = io.BytesIO()
    scan.save(buf, "PNG")
    png = buf.getvalue()
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page(width=560, height=720)
        page.insert_image(fitz.Rect(0, 0, 560, 720), stream=png)  # 整頁掃描影像
        page.insert_text((40, 700), f"Confidential report page {i} footer line bottom")  # 薄文字層
    p = str(tmp_path / "multiscan.pdf")
    doc.save(p)
    doc.close()
    with pytest.raises(ValueError, match="掃描"):
        extract_document(p)


# ───────────────── A. HTML / RTF / TXT / MD ─────────────────

def test_html_file_extraction(tmp_path):
    html = (
        "<html><head><title>T</title></head><body><article>"
        "<h1>區塊鏈新聞標題</h1>"
        "<p>這是一段足夠長的文章內文用來讓主文抽取器判定為正文區塊,"
        "後面再補一段確保字數足夠被 trafilatura 視為文章主體而非樣板。</p>"
        "<p>第二段內容,談論加密貨幣與去中心化金融的最新發展與市場動態,"
        "以及監管環境與機構採用趨勢,確保正文長度充足。</p>"
        "</article></body></html>"
    )
    p = tmp_path / "a.html"
    p.write_text(html, encoding="utf-8")
    text, images = extract_document(str(p))
    assert "區塊鏈新聞標題" in text
    assert images == []


def test_rtf_file_extraction(tmp_path):
    p = tmp_path / "a.rtf"
    p.write_text(r"{\rtf1\ansi\ansicpg950 Hello RTF 內容 World.}", encoding="utf-8")
    text, images = extract_document(str(p))
    assert "Hello RTF" in text and "World" in text and images == []


def test_odt_pure_python_extraction(tmp_path):
    """.odt 純 Python(odfdo,免 LibreOffice):標題/段落/內嵌圖保位/清單/表格。"""
    from odfdo import Document, Frame, Header, Link, List, ListItem, Paragraph, Table

    d = Document("text")
    b = d.body
    b.append(Header(1, "ODT 標題"))
    b.append(Paragraph("ALPHA 前段文字"))
    uri = d.add_file(_png_path(tmp_path))
    pimg = Paragraph("")
    pimg.append(Frame.image_frame(uri, size=("3cm", "3cm")))
    b.append(pimg)
    b.append(Paragraph("BRAVO 後段文字"))
    link_para = Paragraph("前綴 ")
    link_para.append(Link(url="https://blocktempo.ai/", text="動區連結"))
    b.append(link_para)
    lst = List()
    lst.append(ListItem("項目一"))
    lst.append(ListItem("項目二"))
    multi = ListItem()             # 多段落清單項 → 不可被黏成一團
    multi.append(Paragraph("多段一"))
    multi.append(Paragraph("多段二"))
    lst.append(multi)
    b.append(lst)
    t = Table("T")
    t.set_values([["表頭A", "表頭B"], ["值1", "值2"]])
    b.append(t)
    p = str(tmp_path / "doc.odt")
    d.save(p)

    text, images = extract_document(p)
    assert len(images) == 1
    assert text.index("ALPHA 前段文字") < text.index("{{IMG0}}") < text.index("BRAVO 後段文字")
    assert "# ODT 標題" in text
    assert "[動區連結](https://blocktempo.ai/)" in text  # 超連結保真,對齊 DOCX 路徑
    assert "- 項目一" in text and "- 項目二" in text
    assert "- 多段一 多段二" in text and "多段一多段二" not in text  # 多段落不被黏成一團
    assert "| 表頭A | 表頭B |" in text and "| 值1 | 值2 |" in text
    assert "Pictures/" not in text  # 圖框 href 殘影要清乾淨


def test_txt_and_md_passthrough(tmp_path):
    pt = tmp_path / "a.txt"
    pt.write_text("純文字內容", encoding="utf-8")
    assert extract_document(str(pt)) == ("純文字內容", [])
    pm = tmp_path / "a.md"
    pm.write_text("# 標題\n\n內文", encoding="utf-8")
    text, images = extract_document(str(pm))
    assert "# 標題" in text and images == []


# ───────────────── A. 邊界 / 錯誤路徑(明確報錯,不靜默) ─────────────────

def test_legacy_doc_rejected_with_hint(tmp_path):
    f = tmp_path / "old.doc"
    f.write_bytes(b"\xd0\xcf\x11\xe0junk")
    with pytest.raises(ValueError, match="LibreOffice"):
        extract_document(str(f))


def test_image_file_rejected_with_hint(tmp_path):
    f = tmp_path / "scan.png"
    f.write_bytes(_png())
    with pytest.raises(ValueError, match="OCR"):
        extract_document(str(f))


def test_unknown_type_rejected(tmp_path):
    f = tmp_path / "a.xyz"
    f.write_text("x")
    with pytest.raises(ValueError, match="不支援"):
        extract_document(str(f))


# ───────────────── A. markdown → HTML(表格 / 圖 / 程式碼不壞) ─────────────────

def test_md_html_table_renders():
    html = md_to_html("| 欄1 | 欄2 |\n|---|---|\n| a | b |")
    assert "<table>" in html and "<td>a</td>" in html


def test_md_html_image_figure():
    html = md_to_html("文字\n\n![](https://x/a.png)\n\n更多")
    assert 'figure class="article-image"' in html and 'loading="lazy"' in html


def test_md_html_fenced_code():
    html = md_to_html("```\ncode line\n```")
    assert "<code>" in html and "code line" in html


# ───────────────── B. 真實素材(有檔才跑) ─────────────────

def _f(name: str) -> str:
    p = os.path.join(EX, name)
    if not os.path.exists(p):
        pytest.skip(f"缺素材 {name}")
    return p


def test_real_docx_hashkey():
    text, images = extract_document(_f("【新聞稿】_HashKey_PRO將在第二季度正式上線，支援法幣交易對.docx"))
    assert "HashKey" in text and len(images) >= 1 and "{{IMG0}}" in text


def test_real_docx_weex_simplified():
    text, images = extract_document(_f("中文埋点WEEX强势亮相_Consensus_HK，与行业精英携手推动以AI_赋能加密交易.docx"))
    assert len(text) > 300 and len(images) >= 1


def test_real_pdf_text_only():
    text, images = extract_document(_f("Bluefin - article content.pdf"))
    assert "Bluefin" in text and len(images) == 0


def test_real_pdf_tables_and_images():
    """數碼港 PDF:表格→markdown(>=2)、內嵌圖去重成 4 張。"""
    text, images = extract_document(_f("新聞稿數碼港Web3-Innovators-Season系列主題活動.pdf"))
    assert text.count("|---|") >= 2
    assert "活動亮點" in text
    assert len(images) == 4
