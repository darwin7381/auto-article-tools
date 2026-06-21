"""進稿膠合層測試 —— ingest(檔案/URL)+ url_extract。

獨立稽核抓到「URL 路徑 + ingest/_embed_images 零自動化測試」是最高靜默失敗風險,這裡補上:
- 檔案路徑:離線實跑(extract → _embed_images),斷言 {{IMGn}} 換成 markdown 圖、孤兒清掉。
- URL 路徑:mock 網路(httpx/trafilatura/firecrawl/gdoc),斷言派發、fallback、死鏈守門。
"""

from __future__ import annotations

import pytest


# ───────────────────────── _embed_images(離線、純邏輯) ─────────────────────────

def test_embed_images_replace_and_orphan_cleanup(monkeypatch):
    from app.services import ingest as ing

    n = {"i": 0}

    def fake_save(blob, ext):
        n["i"] += 1
        return (f"/p/{n['i']}.{ext}", f"/files/img/{n['i']}.{ext}")

    monkeypatch.setattr(ing, "save_image", fake_save)
    text, urls = ing._embed_images("A{{IMG0}}B{{IMG1}}C{{IMG9}}", [(b"x", "png"), (b"y", "jpeg")])
    assert urls == ["/files/img/1.png", "/files/img/2.jpg"]   # jpeg → jpg
    assert "![](/files/img/1.png)" in text and "![](/files/img/2.jpg)" in text
    assert "{{IMG" not in text                                 # 孤兒 {{IMG9}} 被清掉
    assert text.index("/files/img/1.png") < text.index("/files/img/2.jpg")  # 順序保留


# ───────────────────────── ingest 檔案路徑(離線實跑) ─────────────────────────

async def test_ingest_file_embeds_images(tmp_path, monkeypatch):
    import docx
    from PIL import Image

    from app.services import ingest as ing

    monkeypatch.setattr(ing, "save_image", lambda blob, ext: (f"/p.{ext}", f"/files/x.{ext}"))

    pic = tmp_path / "p.png"
    Image.new("RGB", (10, 10), "red").save(pic)
    d = docx.Document()
    d.add_paragraph("ALPHA 前段")
    d.add_paragraph().add_run().add_picture(str(pic))
    d.add_paragraph("BRAVO 後段")
    fp = tmp_path / "a.docx"
    d.save(str(fp))

    md, urls = await ing.ingest({"file": str(fp)})
    assert urls == ["/files/x.png"]
    assert "![](/files/x.png)" in md and "{{IMG" not in md
    assert md.index("ALPHA") < md.index("![](/files/x.png)") < md.index("BRAVO")  # 圖位保留


async def test_ingest_url_collects_image_urls(monkeypatch):
    from app.services import ingest as ing

    async def fake_extract_url(url):
        return "# 標題\n\n內文一二三 ![](https://x/a.png) 內文 ![alt](https://x/b.jpg) 結尾"

    monkeypatch.setattr(ing, "extract_url", fake_extract_url)
    md, urls = await ing.ingest({"url": "https://example.com/post"})
    assert urls == ["https://x/a.png", "https://x/b.jpg"]
    assert "標題" in md


async def test_ingest_requires_file_or_url():
    from app.services import ingest as ing

    with pytest.raises(ValueError, match="file|url"):
        await ing.ingest({})


async def test_ingest_url_only_collects_http_images(monkeypatch):
    """URL 路徑只收 http(s) 圖;data:/相對路徑不被當圖 URL(記錄此行為)。"""
    from app.services import ingest as ing

    async def fake(url):
        return "文 ![](https://x/a.png) ![](data:image/png;base64,AAAA) ![](/local/b.png)"

    monkeypatch.setattr(ing, "extract_url", fake)
    _, urls = await ing.ingest({"url": "https://e.com"})
    assert urls == ["https://x/a.png"]


def test_extract_html_fragment_fallback(tmp_path):
    """非 article 結構的 HTML 片段 → trafilatura 抽不到主文 → 退回去標籤純文字(不回空)。"""
    from app.services.extract import extract_document

    p = tmp_path / "frag.html"
    p.write_text("<div><span>純片段沒有文章結構</span><b>只有零碎標籤</b></div>", encoding="utf-8")
    text, images = extract_document(str(p))
    assert "純片段沒有文章結構" in text and images == []


# ───────────────────────── url_extract:派發 / 死鏈守門 / fallback ─────────────────────────

async def test_extract_url_dispatch_and_dead_link(monkeypatch):
    from app.services import url_extract as ux

    calls = {"gdoc": 0, "web": 0}

    async def fake_gdoc(doc_id):
        calls["gdoc"] += 1
        return "Google Docs 內容 " * 50  # 夠長

    async def fake_web(url):
        calls["web"] += 1
        return "網頁主文內容 " * 50

    monkeypatch.setattr(ux, "_extract_gdoc", fake_gdoc)
    monkeypatch.setattr(ux, "_extract_webpage", fake_web)

    out = await ux.extract_url("https://docs.google.com/document/d/ABC123/edit")
    assert "Google Docs" in out and calls["gdoc"] == 1 and calls["web"] == 0  # gdoc 走 gdoc 路徑

    out2 = await ux.extract_url("https://medium.com/@x/post")
    assert "網頁主文" in out2 and calls["web"] == 1                            # 一般站走 webpage


async def test_extract_url_dead_link_raises(monkeypatch):
    from app.services import url_extract as ux

    async def fake_web(url):
        return "短"  # < _DEAD_MIN(150)→ 死鏈/殼頁

    monkeypatch.setattr(ux, "_extract_webpage", fake_web)
    with pytest.raises(RuntimeError, match="過短"):
        await ux.extract_url("https://example.com/deleted")


class _FakeResp:
    def __init__(self, text="", status=200):
        self.text = text
        self.status_code = status

    def raise_for_status(self):
        if self.status_code >= 400:
            import httpx

            raise httpx.HTTPStatusError("err", request=None, response=self)  # type: ignore[arg-type]


class _FakeClient:
    def __init__(self, resp):
        self._resp = resp

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return False

    async def get(self, url, headers=None):
        return self._resp


def _article_html(body: str) -> str:
    return f"<html><body><article><h1>標題</h1><p>{body}</p></article></body></html>"


async def test_webpage_direct_trafilatura(monkeypatch):
    from app.services import url_extract as ux

    # 用「不重複」的長內文(repeated 句會被 trafilatura dedup 當樣板濾掉 → 誤觸 fallback)
    long_body = (
        "比特幣在亞洲盤早段走高,市場關注本週的總體經濟數據與聯準會官員談話。"
        "以太坊同步上揚,鏈上活躍地址數較上週明顯回升,DeFi 鎖倉量亦見增加。"
        "分析師指出,機構資金透過現貨 ETF 持續流入,為市場提供結構性支撐。"
        "另一方面,監管面消息分歧,部分司法管轄區對穩定幣發行提出新的資本要求。"
        "交易所數據顯示,衍生品未平倉合約攀升,顯示槓桿需求回溫但波動風險仍在。"
        "整體而言,短線情緒偏多,惟須留意數據公布後的獲利了結賣壓與流動性變化。"
    )

    async def _no_fc(url):
        raise AssertionError("不該走 firecrawl —— 直抓應已成功")

    monkeypatch.setattr(ux, "_firecrawl_html", _no_fc)
    monkeypatch.setattr(ux.httpx, "AsyncClient", lambda **kw: _FakeClient(_FakeResp(_article_html(long_body))))
    md = await ux._extract_webpage("https://news.example.com/post")
    assert "比特幣" in md and "以太坊" in md


async def test_webpage_falls_back_to_firecrawl_when_short(monkeypatch):
    from app.services import url_extract as ux

    # 直抓回殼頁(過短)→ 應改走 firecrawl 渲染後重抽
    monkeypatch.setattr(ux.httpx, "AsyncClient", lambda **kw: _FakeClient(_FakeResp("<html><body>nav</body></html>")))
    fc_used = {"n": 0}
    long_body = "Firecrawl 渲染後抓到的完整繁體中文主文內容,長度足夠被視為正文。" * 8

    async def fake_fc(url):
        fc_used["n"] += 1
        return _article_html(long_body)

    monkeypatch.setattr(ux, "_firecrawl_html", fake_fc)
    md = await ux._extract_webpage("https://spa.example.com/post")
    assert fc_used["n"] == 1 and ("標題" in md or "主文內容" in md)  # 確實走了 fallback


async def test_webpage_404_does_not_fallback(monkeypatch):
    from app.services import url_extract as ux

    monkeypatch.setattr(ux.httpx, "AsyncClient", lambda **kw: _FakeClient(_FakeResp("", status=404)))

    async def fake_fc(url):
        raise AssertionError("404 不該走 firecrawl fallback")

    monkeypatch.setattr(ux, "_firecrawl_html", fake_fc)
    with pytest.raises(Exception):  # noqa: B017,PT011  404 直接往上拋,不浪費 fallback
        await ux._extract_webpage("https://example.com/missing")
